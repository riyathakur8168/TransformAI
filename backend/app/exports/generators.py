import io
import json
import os
from typing import Dict, Any

# PDF ReportLab imports
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
from reportlab.lib import colors

# DOCX imports
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

# PPTX imports
import pptx
from pptx.util import Inches as PPTInches, Pt as PPTPt
from pptx.dml.color import RGBColor as PPTRGBColor

def generate_pdf_bytes(title: str, content: str, output_type: str) -> bytes:
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        rightMargin=40, leftMargin=40, topMargin=40, bottomMargin=40
    )
    styles = getSampleStyleSheet()

    # Custom styles
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontSize=20,
        leading=24,
        textColor=colors.HexColor('#0F172A'),
        spaceAfter=12
    )
    h2_style = ParagraphStyle(
        'Heading2Custom',
        parent=styles['Heading2'],
        fontSize=14,
        leading=18,
        textColor=colors.HexColor('#0284C7'),
        spaceBefore=10,
        spaceAfter=6
    )
    body_style = ParagraphStyle(
        'BodyCustom',
        parent=styles['Normal'],
        fontSize=10.5,
        leading=15,
        textColor=colors.HexColor('#334155'),
        spaceAfter=8
    )

    story = []
    story.append(Paragraph(f"<b>TransformAI Document Export</b>", ParagraphStyle('SubHeader', parent=body_style, textColor=colors.HexColor('#64748B'), fontSize=9)))
    story.append(Spacer(1, 4))
    story.append(Paragraph(title, title_style))
    story.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor('#0284C7'), spaceAfter=15))

    if output_type == "presentation" and content.strip().startswith("["):
        try:
            slides = json.loads(content)
            for s in slides:
                story.append(Paragraph(f"Slide {s.get('slide', 1)}: {s.get('title', '')}", h2_style))
                if s.get("subtitle"):
                    story.append(Paragraph(f"<i>{s.get('subtitle')}</i>", body_style))
                for b in s.get("bullets", []):
                    story.append(Paragraph(f"• {b}", body_style))
                story.append(Spacer(1, 10))
        except Exception:
            story.append(Paragraph(content.replace("\n", "<br/>"), body_style))
    else:
        # Convert markdown formatted headers & bullet lists into ReportLab Paragraphs
        for line in content.splitlines():
            line_str = line.strip()
            if not line_str:
                story.append(Spacer(1, 4))
            elif line_str.startswith("# "):
                story.append(Paragraph(line_str[2:], title_style))
            elif line_str.startswith("## ") or line_str.startswith("### "):
                story.append(Paragraph(line_str.lstrip("#").strip(), h2_style))
            elif line_str.startswith("* ") or line_str.startswith("- ") or line_str.startswith("• "):
                story.append(Paragraph(f"• {line_str[2:]}", body_style))
            else:
                story.append(Paragraph(line_str, body_style))

    doc.build(story)
    return buffer.getvalue()


def generate_docx_bytes(title: str, content: str, output_type: str) -> bytes:
    doc = docx.Document()
    
    # Title
    title_p = doc.add_paragraph()
    run = title_p.add_run(title)
    run.font.size = Pt(22)
    run.font.bold = True
    run.font.color.rgb = RGBColor(15, 23, 42)
    
    # Subtitle
    sub_p = doc.add_paragraph("Exported from TransformAI Engine | Grounded Content Output")
    sub_p.runs[0].font.italic = True
    sub_p.runs[0].font.color.rgb = RGBColor(100, 116, 139)
    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    if output_type == "presentation" and content.strip().startswith("["):
        try:
            slides = json.loads(content)
            for s in slides:
                doc.add_heading(f"Slide {s.get('slide', 1)}: {s.get('title', '')}", level=2)
                if s.get("subtitle"):
                    p = doc.add_paragraph(s.get("subtitle"))
                    p.runs[0].font.italic = True
                for b in s.get("bullets", []):
                    doc.add_paragraph(b, style='List Bullet')
        except Exception:
            doc.add_paragraph(content)
    else:
        for line in content.splitlines():
            line_str = line.strip()
            if not line_str:
                continue
            if line_str.startswith("# "):
                doc.add_heading(line_str[2:], level=1)
            elif line_str.startswith("## "):
                doc.add_heading(line_str[3:], level=2)
            elif line_str.startswith("### "):
                doc.add_heading(line_str[4:], level=3)
            elif line_str.startswith("* ") or line_str.startswith("- ") or line_str.startswith("• "):
                doc.add_paragraph(line_str[2:], style='List Bullet')
            else:
                doc.add_paragraph(line_str)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def generate_pptx_bytes(title: str, content: str) -> bytes:
    prs = pptx.Presentation()

    # Parse slides if available or build from content
    slides_data = []
    if content.strip().startswith("["):
        try:
            slides_data = json.loads(content)
        except Exception:
            pass

    if not slides_data:
        # Fallback slides from text
        slides_data = [
            {"slide": 1, "title": title, "subtitle": "TransformAI Grounded Presentation Deck", "bullets": ["Generated Content Overview"]},
            {"slide": 2, "title": "Main Content Details", "subtitle": "Key Analysis", "bullets": [line[:100] for line in content.splitlines() if line.strip() and not line.startswith("#")][:4]}
        ]

    for s in slides_data:
        blank_slide_layout = prs.slide_layouts[6]
        slide = prs.slides.add_slide(blank_slide_layout)

        # Title Box
        txBox = slide.shapes.add_textbox(PPTInches(0.8), PPTInches(0.6), PPTInches(8.4), PPTInches(1.2))
        tf = txBox.text_frame
        p = tf.paragraphs[0]
        p.text = s.get("title", title)
        p.font.size = PPTPt(28)
        p.font.bold = True
        p.font.color.rgb = PPTRGBColor(15, 23, 42)

        if s.get("subtitle"):
            p2 = tf.add_paragraph()
            p2.text = s.get("subtitle", "")
            p2.font.size = PPTPt(14)
            p2.font.color.rgb = PPTRGBColor(2, 132, 199)

        # Bullets Box
        txBoxBody = slide.shapes.add_textbox(PPTInches(0.8), PPTInches(2.0), PPTInches(8.4), PPTInches(4.8))
        tfBody = txBoxBody.text_frame
        tfBody.word_wrap = True

        bullets = s.get("bullets", [])
        for i, b in enumerate(bullets):
            p_bullet = tfBody.paragraphs[0] if i == 0 else tfBody.add_paragraph()
            p_bullet.text = f"• {b}"
            p_bullet.font.size = PPTPt(18)
            p_bullet.font.color.rgb = PPTRGBColor(51, 65, 85)
            p_bullet.space_after = PPTPt(14)

    buffer = io.BytesIO()
    prs.save(buffer)
    return buffer.getvalue()


def generate_export_file(title: str, content: str, output_type: str, export_format: str) -> bytes:
    fmt = export_format.lower()
    if fmt == "pdf":
        return generate_pdf_bytes(title, content, output_type)
    elif fmt == "docx":
        return generate_docx_bytes(title, content, output_type)
    elif fmt == "pptx":
        return generate_pptx_bytes(title, content)
    else:  # TXT
        text_out = f"TITLE: {title}\nEXPORTED FROM: TransformAI Platform\n\n{content}"
        return text_out.encode("utf-8")
